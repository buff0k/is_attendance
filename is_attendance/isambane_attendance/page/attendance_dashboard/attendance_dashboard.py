# Copyright (c) 2026, BuFf0k and contributors
# For license information, please see license.txt

"""
Attendance Dashboard - a client-side page (attendance_dashboard.js) that
drives the "Attendance Compliance Summary" report via
frappe.desk.query_report.run for the on-screen table, and this module's
export_compliance_excel for the download.

The export has two kinds of sheet:

- **One sheet per employee**: a day-by-day breakdown (Date, Day, Day Type,
  In, Out, Hours Worked, Public Holiday, On Leave, and the same
  Missed/No Out/No In/Late In/Early Out flags, per day) - the raw
  Date/In/Out values come from attendance_compliance_summary.compute(),
  but everything derivable purely from what's on the sheet is a **live
  Excel formula**, not a Python-computed value - Day/Day Type read off
  that row's own Date cell (=TEXT(...,"dddd") / WEEKDAY-based CHOOSE), and
  Hours Worked reads off that row's own In/Out cells ((Out-In)*24 - see
  _write_daily_sheet). Date is a real native Excel date value and In/Out
  are real native Excel TIME values (never a datetime merely formatted to
  look like one) - none of these three are text or Python-computed, so
  editing a Date, In, or Out cell by hand ripples through correctly.
  Deliberately separate from - and allowed to disagree with -
  attendance_compliance_summary.compute()'s own Python-computed
  hours_worked (which sums every clock-in/out pair that day, used for
  Missed/Late In/Early Out and the on-screen Dashboard): that's report
  logic answering "was this day compliant," this sheet's Hours Worked
  column is spreadsheet logic answering "what does Out minus In say,"
  and the two only have to agree on a day with a single session.
- **Report**: a compact per-employee summary, grouped by Weekday/Saturday/
  Sunday under merged group headers rather than one flat 22-column table.
  Its metric cells are **live Excel formulas** (COUNTIF/COUNTIFS) reading
  straight from that employee's own sheet - not pre-computed Python values
  - matching how the reference workbook's Report sheet actually worked, and
  meaning the numbers stay correct even if someone edits a day's flags by
  hand directly in their sheet afterwards.
"""

from __future__ import annotations

import re
from io import BytesIO

import frappe
from frappe import _
from frappe.utils import now_datetime

from is_attendance.isambane_attendance.report.attendance_compliance_summary.attendance_compliance_summary import (
	DAY_TYPES,
	METRICS,
	compute,
)
from is_attendance.permissions import responsible_branches_for_user

# Matches the "Attendance Compliance Summary" Report doctype's own roles -
# that gate only applies when going through the query-report UI/API, not to
# this standalone whitelisted export, so it's re-checked explicitly here.
EXPORT_ROLES = {"System Manager", "HR Manager", "HR User", "Payroll Manager", "Payroll User"}

# Fixed column order for every per-employee daily sheet - the Report
# sheet's formulas hardcode letters derived from this same list, so the
# two can never drift apart.
DAILY_SHEET_COLUMNS = [
	("date", "Date"),
	("day", "Day"),
	("day_type", "Day Type"),
	("in_time", "In"),
	("out_time", "Out"),
	("hours_worked", "Hours Worked"),
	("public_holiday", "Public Holiday"),
	("on_leave", "On Leave"),
	("half_day_leave", "Half Day Leave"),
	("leave_status", "Leave Status"),
	("missed", "Missed"),
	("no_out", "No Out"),
	("no_in", "No In"),
	("late_in", "Late In"),
	("early_out", "Early Out"),
]
DAILY_SHEET_DATA_START_ROW = 3  # row 1 = title, row 2 = header

METRIC_LABELS = dict(METRICS)  # {"missed": "Missed", "no_out": "No Out", ...}

INVALID_SHEET_NAME_CHARS = re.compile(r"[:\\/?*\[\]]")


@frappe.whitelist()
def export_compliance_excel(filters=None):
	if not EXPORT_ROLES & set(frappe.get_roles()):
		frappe.throw(_("You are not permitted to export this report."), frappe.PermissionError)

	filters = frappe.parse_json(filters) if isinstance(filters, str) else (filters or {})

	summary_rows, daily_detail = compute(filters)

	workbook = _build_workbook(summary_rows, daily_detail)

	buffer = BytesIO()
	workbook.save(buffer)

	frappe.response["filename"] = f"attendance_compliance_{now_datetime().strftime('%Y%m%d_%H%M%S')}.xlsx"
	frappe.response["filecontent"] = buffer.getvalue()
	# "download" (not "binary") so the Content-Type gets guessed from the
	# .xlsx filename (frappe/utils/response.py::as_raw) rather than a
	# generic application/octet-stream.
	frappe.response["type"] = "download"


@frappe.whitelist()
def get_daily_detail(filters=None, employee=None):
	"""Drill-down data for one employee row on the Dashboard's table -
	reuses compute() directly (same call export_compliance_excel makes) so
	this can never disagree with the summary row it expands from.

	employee is checked against responsible_branches_for_user() explicitly
	because it's a direct argument, not something that went through
	compute()'s/_resolve_employees()'s own filtering chain - without this a
	branch-restricted user could ask for another branch's employee by name
	even though they can't see them in the list this expands from."""
	if not EXPORT_ROLES & set(frappe.get_roles()):
		frappe.throw(_("You are not permitted to view this data."), frappe.PermissionError)

	if not employee:
		frappe.throw(_("Employee is required."))

	branches = responsible_branches_for_user()
	if branches:
		employee_branch = frappe.db.get_value("Employee", employee, "branch")
		if employee_branch not in branches:
			frappe.throw(_("Not permitted to view this employee's data."), frappe.PermissionError)

	filters = frappe.parse_json(filters) if isinstance(filters, str) else (filters or {})

	_summary_rows, daily_detail = compute(filters)
	return daily_detail.get(employee, [])


TIME_FIELDS = {"in_time", "out_time"}
DATE_FIELDS = {"date"}
# Which category Excel's own Format Cells dialog shows (Date/Time vs
# Custom) is decided by matching the cell's literal format code string
# against Excel's own canonical list for the viewer's locale - "mm-dd-yy"
# (openpyxl's builtin id 14 string) matched and shows as a real Date cell,
# confirmed on a real export. For Time, id 20's string ("h:mm", no
# seconds) did not match, and neither did a plain "hh:mm" guess - the
# format this site's own Excel actually applies when you pick "Time" >
# 24-hour from Format Cells includes seconds ("13:30:55"), which is
# exactly openpyxl's OTHER genuine Time builtin, id 21 ("h:mm:ss") - the
# one 24-hour Time builtin with seconds; only id 20 (no seconds) was tried
# before. Showing seconds is a closer match to the underlying data anyway
# - the raw device checkin timestamps this comes from are second-precise.
DATE_NUMBER_FORMAT = "mm-dd-yy"
TIME_NUMBER_FORMAT = "h:mm:ss"


def _build_workbook(summary_rows: list[dict], daily_detail: dict[str, list[dict]]):
	from openpyxl import Workbook

	workbook = Workbook()
	workbook.remove(workbook.active)

	employee_names = {row["employee"]: row.get("employee_name") for row in summary_rows}
	sheet_names: dict[str, str] = {}
	used_sheet_names: set[str] = set()

	for employee, day_rows in daily_detail.items():
		sheet_name = _unique_sheet_name(employee, employee_names.get(employee), used_sheet_names)
		used_sheet_names.add(sheet_name)
		sheet_names[employee] = sheet_name
		_write_daily_sheet(workbook.create_sheet(sheet_name), sheet_name, day_rows)

	# Report goes first (Excel opens to whichever sheet was active last -
	# create it last, then move it to the front and select it).
	report_sheet = workbook.create_sheet("Report", 0)
	_write_report_sheet(report_sheet, summary_rows, daily_detail, sheet_names)
	workbook.active = 0

	return workbook


def _write_daily_sheet(sheet, title: str, day_rows: list[dict]):
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter

	thin = Side(style="thin", color="000000")
	border = Border(left=thin, right=thin, top=thin, bottom=thin)
	title_font = Font(bold=True, size=13)
	header_font = Font(bold=True, color="FFFFFF")
	header_fill = PatternFill("solid", fgColor="305496")
	header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
	center_align = Alignment(horizontal="center", vertical="center")

	column_count = len(DAILY_SHEET_COLUMNS)

	sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
	title_cell = sheet.cell(row=1, column=1, value=title)
	title_cell.font = title_font
	title_cell.alignment = center_align
	for col_index in range(1, column_count + 1):
		sheet.cell(row=1, column=col_index).border = border
	sheet.row_dimensions[1].height = 20

	for col_index, (_fieldname, label) in enumerate(DAILY_SHEET_COLUMNS, start=1):
		cell = sheet.cell(row=2, column=col_index, value=label)
		cell.font = header_font
		cell.fill = header_fill
		cell.alignment = header_align
		cell.border = border
	sheet.row_dimensions[2].height = 26

	date_col = _daily_column_letter("date")
	in_col = _daily_column_letter("in_time")
	out_col = _daily_column_letter("out_time")

	for row_offset, row in enumerate(day_rows):
		row_index = DAILY_SHEET_DATA_START_ROW + row_offset
		date_ref = f"{date_col}{row_index}"
		in_ref = f"{in_col}{row_index}"
		out_ref = f"{out_col}{row_index}"

		for col_index, (fieldname, _label) in enumerate(DAILY_SHEET_COLUMNS, start=1):
			if fieldname == "day":
				# Live formula off the Date cell, not a precomputed string -
				# stays correct if someone edits a row's own Date by hand,
				# same "never quietly drift from the sheet's own data" reason
				# the Report sheet's metric cells are formulas (see module
				# docstring).
				value = f'=TEXT({date_ref},"dddd")'
			elif fieldname == "day_type":
				# WEEKDAY(date, 2) returns 1=Monday...7=Sunday - CHOOSE maps
				# that straight onto the same Weekday/Saturday/Sunday
				# classification _day_type() uses server-side, purely from
				# the Date cell, so it can never disagree with it.
				value = f'=CHOOSE(WEEKDAY({date_ref},2),"Weekday","Weekday","Weekday","Weekday","Weekday","Saturday","Sunday")'
			elif fieldname in TIME_FIELDS:
				# Pure time-of-day (no date component at all) - a real Excel
				# TIME value, not a full datetime formatted to merely look
				# like one, so the Hours Worked formula below (and anything
				# else referencing these cells) is doing genuine time
				# arithmetic on them.
				stamp = row.get(fieldname)
				value = stamp.time() if stamp else None
			elif fieldname == "hours_worked":
				# A real Excel formula off this row's own In/Out cells - the
				# report's own Python-computed hours_worked (which sums
				# every clock-in/out pair that day, not just the first and
				# last) is what drives Missed/Late In/Early Out and the
				# on-screen Dashboard; this column is a separate, simpler,
				# self-contained spreadsheet figure computed purely from
				# what's shown here, same as anyone doing =Out-In by hand
				# would get.
				value = f"=IF(AND({in_ref}<>\"\",{out_ref}<>\"\"),({out_ref}-{in_ref})*24,0)"
			else:
				value = _excel_value(row.get(fieldname))

			cell = sheet.cell(row=row_index, column=col_index, value=value)
			cell.border = border
			if fieldname in DATE_FIELDS:
				cell.number_format = DATE_NUMBER_FORMAT
			elif fieldname in TIME_FIELDS:
				cell.number_format = TIME_NUMBER_FORMAT
			elif fieldname == "hours_worked":
				cell.number_format = "0.00"
			if fieldname not in ("day", "day_type", "public_holiday"):
				cell.alignment = center_align

	widths = {
		"date": 11,
		"day": 10,
		"day_type": 9,
		"in_time": 8,
		"out_time": 8,
		"hours_worked": 11,
		"public_holiday": 20,
		"on_leave": 9,
		"half_day_leave": 14,
		"leave_status": 12,
		"missed": 8,
		"no_out": 8,
		"no_in": 8,
		"late_in": 8,
		"early_out": 9,
	}
	for col_index, (fieldname, _label) in enumerate(DAILY_SHEET_COLUMNS, start=1):
		sheet.column_dimensions[get_column_letter(col_index)].width = widths.get(fieldname, 10)

	sheet.freeze_panes = "A3"


def _write_report_sheet(
	sheet,
	summary_rows: list[dict],
	daily_detail: dict[str, list[dict]],
	sheet_names: dict[str, str],
):
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter

	thin = Side(style="thin", color="000000")
	border = Border(left=thin, right=thin, top=thin, bottom=thin)
	title_font = Font(bold=True, size=14)
	group_font = Font(bold=True, color="FFFFFF")
	group_fill = PatternFill("solid", fgColor="1F3864")
	header_font = Font(bold=True, color="FFFFFF")
	header_fill = PatternFill("solid", fgColor="305496")
	header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
	center_align = Alignment(horizontal="center", vertical="center")

	identity_columns = [("employee", "Employee"), ("employee_name", "Employee Name"), ("branch", "Branch")]
	metric_fieldnames = ["total"] + [metric for metric, _label in METRICS]
	metric_labels = ["Total"] + [METRIC_LABELS[metric] for metric in metric_fieldnames[1:]]

	column_count = len(identity_columns) + len(DAY_TYPES) * len(metric_fieldnames)

	# Title row
	sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
	title_cell = sheet.cell(row=1, column=1, value="Isambane Attendance Compliance")
	title_cell.font = title_font
	title_cell.alignment = center_align
	for col_index in range(1, column_count + 1):
		sheet.cell(row=1, column=col_index).border = border
	sheet.row_dimensions[1].height = 22

	# Identity columns span both header rows (no day-type grouping above them)
	col = 1
	for fieldname, label in identity_columns:
		sheet.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
		cell = sheet.cell(row=2, column=col, value=label)
		cell.font = header_font
		cell.fill = header_fill
		cell.alignment = header_align
		for r in (2, 3):
			sheet.cell(row=r, column=col).border = border
		col += 1

	# One merged group header per day type, spanning its metric columns,
	# with the individual metric names on the row below.
	day_type_start_col = {}
	for day_type in DAY_TYPES:
		day_type_start_col[day_type] = col
		end_col = col + len(metric_fieldnames) - 1
		sheet.merge_cells(start_row=2, start_column=col, end_row=2, end_column=end_col)
		group_cell = sheet.cell(row=2, column=col, value=day_type)
		group_cell.font = group_font
		group_cell.fill = group_fill
		group_cell.alignment = center_align
		for c in range(col, end_col + 1):
			sheet.cell(row=2, column=c).border = border

		for offset, label in enumerate(metric_labels):
			cell = sheet.cell(row=3, column=col + offset, value=label)
			cell.font = header_font
			cell.fill = header_fill
			cell.alignment = header_align
			cell.border = border

		col = end_col + 1

	sheet.row_dimensions[2].height = 20
	sheet.row_dimensions[3].height = 26

	# Data rows - identity columns are plain values, every metric column is
	# a live formula reading that employee's own daily sheet.
	data_start_row = 4
	for row_offset, row in enumerate(summary_rows):
		row_index = data_start_row + row_offset
		employee = row["employee"]
		day_rows = daily_detail.get(employee, [])
		last_row = DAILY_SHEET_DATA_START_ROW + len(day_rows) - 1
		sheet_ref = _formula_sheet_ref(sheet_names[employee])
		day_type_col = _daily_column_letter("day_type")

		col = 1
		for fieldname, _label in identity_columns:
			cell = sheet.cell(row=row_index, column=col, value=row.get(fieldname))
			cell.border = border
			if fieldname != "employee_name":
				cell.alignment = center_align
			col += 1

		for day_type in DAY_TYPES:
			day_range = f"{sheet_ref}!${day_type_col}${DAILY_SHEET_DATA_START_ROW}:${day_type_col}${last_row}"

			total_cell = sheet.cell(row=row_index, column=col, value=f'=COUNTIF({day_range},"{day_type}")')
			total_cell.border = border
			total_cell.alignment = center_align
			col += 1

			for metric, _label in METRICS:
				metric_col = _daily_column_letter(metric)
				metric_range = (
					f"{sheet_ref}!${metric_col}${DAILY_SHEET_DATA_START_ROW}:${metric_col}${last_row}"
				)
				formula = f'=COUNTIFS({day_range},"{day_type}",{metric_range},"Yes")'
				cell = sheet.cell(row=row_index, column=col, value=formula)
				cell.border = border
				cell.alignment = center_align
				col += 1

	# Column widths - identity columns wider, metric columns narrow now
	# that they're no longer prefixed with the day-type name.
	sheet.column_dimensions["A"].width = 12
	sheet.column_dimensions["B"].width = 22
	sheet.column_dimensions["C"].width = 14
	metric_width = {"Total": 8, "Missed": 9, "No Out": 9, "No In": 8, "Late In": 9, "Early Out": 10}
	col = 4
	for _day_type in DAY_TYPES:
		for label in metric_labels:
			sheet.column_dimensions[get_column_letter(col)].width = metric_width.get(label, 9)
			col += 1

	sheet.freeze_panes = f"D{data_start_row}"


def _daily_column_letter(fieldname: str) -> str:
	from openpyxl.utils import get_column_letter

	index = next(i for i, (name, _label) in enumerate(DAILY_SHEET_COLUMNS, start=1) if name == fieldname)
	return get_column_letter(index)


def _formula_sheet_ref(sheet_name: str) -> str:
	return "'" + sheet_name.replace("'", "''") + "'"


def _excel_value(value):
	if isinstance(value, bool):
		return "Yes" if value else ""
	return value


def _unique_sheet_name(employee: str, employee_name: str | None, used_names: set[str]) -> str:
	label = f"{employee} {employee_name}".strip() if employee_name else employee
	label = INVALID_SHEET_NAME_CHARS.sub(" ", label).strip() or employee
	base = label[:31]

	if base not in used_names:
		return base

	# Truncated names can collide (two employees whose first 31 characters
	# match) - append a short numeric suffix, still within the 31-char
	# limit, until it's unique.
	for suffix in range(2, 1000):
		marker = f" ({suffix})"
		candidate = base[: 31 - len(marker)] + marker
		if candidate not in used_names:
			return candidate

	frappe.throw(_("Could not generate a unique Excel sheet name for Employee {0}.").format(employee))
